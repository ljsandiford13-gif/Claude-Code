"""Complete the Final_Draft_HSFB_Workbook layout with live formulas, anonymised stores and charts."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment

from data2 import ROWS, STORE_A, STORE_B, CATEGORIES

OUT = sys.argv[1]
NAVY, BLUE, ORANGE, AQUA = "184F95", "2A78D6", "EB6834", "1BAF7A"
PALE_BLUE, PALE_RED, LIGHT_GREY, MID_GREY = "E3EEFB", "FBE3E3", "F2F2F2", "D9D9D9"
FONT = "Arial"
f_base = Font(name=FONT, size=10); f_bold = Font(name=FONT, size=10, bold=True)
f_input = Font(name=FONT, size=10, color="0000FF"); f_hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color=NAVY); f_note = Font(name=FONT, size=9, italic=True, color="595959")
fill_hdr = PatternFill("solid", fgColor=NAVY); fill_grey = PatternFill("solid", fgColor=LIGHT_GREY)
thin = Side(style="thin", color=MID_GREY); border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")
PCT = '0.0%;-0.0%;0.0%'; MONEY = '#,##0.00'

wb = Workbook()

# ------------------------------------------------------------------ Comparison Data
ws = wb.active; ws.title = "Comparison Data"
headers = ["Category", "Food First Product", "Food First Basis", "Food First Std Price (BBD)",
           f"{STORE_A}: Quantity/Basis", f"{STORE_A} Unit", f"{STORE_A} Shelf Price (BBD)", f"{STORE_A} Std Price (BBD)",
           f"{STORE_B}: Quantity/Basis", f"{STORE_B} Unit", f"{STORE_B} Shelf Price (BBD)", f"{STORE_B} Std Price (BBD)",
           f"Food First vs {STORE_A} %", f"Food First vs {STORE_B} %", f"{STORE_A} vs {STORE_B} %", "Cheapest", "Note"]
for i, h in enumerate(headers, 1):
    c = ws.cell(1, i, h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = border
ws.row_dimensions[1].height = 42
for i, w in enumerate([12, 26, 10, 12, 16, 11, 12, 12, 16, 11, 12, 12, 13, 13, 13, 15, 50], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
FIRST, LAST = 2, 1 + len(ROWS)
for i, r in enumerate(ROWS):
    rr = FIRST + i
    cat, item, basis, ff, aq, au, ash, a, bq, bu, bsh, b, note = r
    for ci, v in enumerate([cat, item, basis, ff, aq, au, ash, a, bq, bu, bsh, b], 1):
        c = ws.cell(rr, ci, v); c.border = border; c.font = f_base
    for ci in (4, 7, 8, 11, 12):
        ws.cell(rr, ci).font = f_input; ws.cell(rr, ci).number_format = MONEY
    ws.cell(rr, 13, f"=D{rr}/H{rr}-1"); ws.cell(rr, 14, f"=D{rr}/L{rr}-1"); ws.cell(rr, 15, f"=H{rr}/L{rr}-1")
    ws.cell(rr, 16, f'=IF(D{rr}<=MIN(H{rr},L{rr}),"Food First",IF(H{rr}<=L{rr},"{STORE_A}","{STORE_B}"))')
    for ci in (13, 14, 15):
        c = ws.cell(rr, ci); c.number_format = PCT; c.alignment = right; c.border = border; c.font = f_base
    ws.cell(rr, 16).border = border; ws.cell(rr, 16).font = f_base; ws.cell(rr, 16).alignment = center
    c = ws.cell(rr, 17, note); c.font = f_note; c.border = border
    if ash != a:
        ws.cell(rr, 8).comment = Comment(f"Standardised from shelf price {ash} ({aq}). {note}", "Analyst")
    if bsh != b:
        ws.cell(rr, 12).comment = Comment(f"Standardised from shelf price {bsh} ({bq}). {note}", "Analyst")
TOT = LAST + 2
ws.cell(TOT, 2, "Matched-basket total (sum of standardised prices)").font = f_bold
for ci, col in ((4, "D"), (8, "H"), (12, "L")):
    c = ws.cell(TOT, ci, f"=SUM({col}{FIRST}:{col}{LAST})"); c.font = f_bold; c.number_format = MONEY; c.border = border
    c.fill = PatternFill("solid", fgColor="FFF7CC")
ws.cell(TOT + 2, 1, "Legend").font = f_bold
ws.cell(TOT + 3, 1, "Blue text = recorded price (editable input).  Black = formula.  % difference = first-named store / second-named store - 1; negative = first-named store cheaper.").font = f_note
ws.cell(TOT + 4, 1, f"Cheapest = lowest standardised price of the three (ties go to Food First, then {STORE_A}).  Std price = shelf price converted to the Food First basis; hover a commented cell for the conversion.").font = f_note
ws.cell(TOT + 5, 1, f"Stores are anonymised per HSFB's undertaking to the retailers: {STORE_A} and {STORE_B}.").font = f_note
ws.conditional_formatting.add(f"M{FIRST}:O{LAST}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_BLUE)))
ws.conditional_formatting.add(f"M{FIRST}:O{LAST}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_RED)))
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:Q{LAST}"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "1:1"

R = lambda col: f"'Comparison Data'!${col}${FIRST}:${col}${LAST}"
CAT, FFA, FFB, AB, CHEAP, D_, H_, L_ = R("A"), R("M"), R("N"), R("O"), R("P"), R("D"), R("H"), R("L")

# ------------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary")
sm.column_dimensions["A"].width = 48; sm.column_dimensions["B"].width = 16
for i, h in enumerate(["Metric", "Value"], 1):
    c = sm.cell(1, i, h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = border
metrics = [
    ("Total items compared", f"=COUNTA({CAT})", "0"),
    (f"Food First cheaper than {STORE_A}", f'=COUNTIF({FFA},"<0")', "0"),
    (f"Food First more expensive than {STORE_A}", f'=COUNTIF({FFA},">0")', "0"),
    (f"Food First cheaper than {STORE_B}", f'=COUNTIF({FFB},"<0")', "0"),
    (f"Food First more expensive than {STORE_B}", f'=COUNTIF({FFB},">0")', "0"),
    ("Food First cheapest of all three stores", f'=COUNTIF({CHEAP},"Food First")', "0"),
    (f"Mean % difference, Food First vs {STORE_A}", f"=AVERAGE({FFA})", PCT),
    (f"Median % difference, Food First vs {STORE_A}", f"=MEDIAN({FFA})", PCT),
    (f"Mean % difference, Food First vs {STORE_B}", f"=AVERAGE({FFB})", PCT),
    (f"Median % difference, Food First vs {STORE_B}", f"=MEDIAN({FFB})", PCT),
    (f"Mean % difference, {STORE_A} vs {STORE_B}", f"=AVERAGE({AB})", PCT),
    ("Matched-basket total, Food First (BBD)", f"=SUM({D_})", MONEY),
    (f"Matched-basket total, {STORE_A} (BBD)", f"=SUM({H_})", MONEY),
    (f"Matched-basket total, {STORE_B} (BBD)", f"=SUM({L_})", MONEY),
    (f"Basket-total difference, Food First vs {STORE_A}", "=B13/B14-1", PCT),
    (f"Basket-total difference, Food First vs {STORE_B}", "=B13/B15-1", PCT),
]
for i, (label, f, fmt) in enumerate(metrics):
    rr = 2 + i
    sm.cell(rr, 1, label).font = f_base; sm.cell(rr, 1).border = border
    c = sm.cell(rr, 2, f); c.number_format = fmt; c.font = f_base; c.alignment = right; c.border = border
# extra block: cheapest store and same-price counts
X0 = 2 + len(metrics) + 1
for i, h in enumerate(["Cheapest of the three stores", "Items"], 1):
    c = sm.cell(X0, i, h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = border
for i, s in enumerate(["Food First", STORE_A, STORE_B]):
    sm.cell(X0 + 1 + i, 1, s).font = f_base; sm.cell(X0 + 1 + i, 1).border = border
    c = sm.cell(X0 + 1 + i, 2, f'=COUNTIF({CHEAP},A{X0+1+i})'); c.number_format = "0"; c.font = f_base; c.alignment = right; c.border = border
sm.cell(X0 + 5, 1, f"Same price: Food First = {STORE_B}").font = f_base
sm.cell(X0 + 5, 2, f'=COUNTIF({FFB},0)').number_format = "0"
sm.cell(X0 + 6, 1, f"Median % difference, {STORE_A} vs {STORE_B}").font = f_base
sm.cell(X0 + 6, 2, f"=MEDIAN({AB})").number_format = PCT
sm.cell(X0 + 7, 1, f"{STORE_A} cheaper than {STORE_B} (items)").font = f_base
sm.cell(X0 + 7, 2, f'=COUNTIF({AB},"<0")').number_format = "0"
sm.cell(X0 + 8, 1, f"{STORE_B} cheaper than {STORE_A} (items)").font = f_base
sm.cell(X0 + 8, 2, f'=COUNTIF({AB},">0")').number_format = "0"
for rr in range(X0 + 5, X0 + 9):
    for ci in (1, 2): sm.cell(rr, ci).border = border; sm.cell(rr, ci).font = f_base
    sm.cell(rr, 2).alignment = right
sm.cell(X0 + 10, 1, "All values are formulas over the Comparison Data tab. Negative % = Food First (or the first-named store) cheaper.").font = f_note
for rng in ("B8:B12", "B16:B17"):
    sm.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_BLUE)))
    sm.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_RED)))

# ------------------------------------------------------------------ Category Analysis
ca = wb.create_sheet("Category Analysis")
hdrs = ["Category", "Items listed", "Items matched", f"Mean FF vs {STORE_A}", f"Median FF vs {STORE_A}", f"Mean FF vs {STORE_B}", f"Median FF vs {STORE_B}",
        f"Mean {STORE_A} vs {STORE_B}", f"FF cheaper vs {STORE_A}", f"FF cheaper vs {STORE_B}", "FF cheaper than both", "FF cheapest of three"]
for i, h in enumerate(hdrs, 1):
    c = ca.cell(1, i, h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = border
ca.row_dimensions[1].height = 44
for i, w in enumerate([16, 10, 10, 13, 13, 13, 13, 13, 13, 13, 13, 13], 1):
    ca.column_dimensions[get_column_letter(i)].width = w
for i, cat in enumerate(CATEGORIES + ["All"]):
    rr = 2 + i; is_all = cat == "All"
    ca.cell(rr, 1, cat).font = f_bold if is_all else f_base
    if is_all:
        f = [f"=COUNTA({CAT})", f"=COUNT({FFA})", f"=AVERAGE({FFA})", f"=MEDIAN({FFA})", f"=AVERAGE({FFB})", f"=MEDIAN({FFB})", f"=AVERAGE({AB})",
             f'=COUNTIF({FFA},"<0")', f'=COUNTIF({FFB},"<0")', f'=SUMPRODUCT(({FFA}<0)*({FFB}<0))', f'=COUNTIF({CHEAP},"Food First")']
    else:
        f = [f"=COUNTIF({CAT},$A{rr})", f"=COUNTIFS({CAT},$A{rr},{FFA},\"<>\")",
             f"=AVERAGEIFS({FFA},{CAT},$A{rr})", None, f"=AVERAGEIFS({FFB},{CAT},$A{rr})", None, f"=AVERAGEIFS({AB},{CAT},$A{rr})",
             f'=COUNTIFS({CAT},$A{rr},{FFA},"<0")', f'=COUNTIFS({CAT},$A{rr},{FFB},"<0")',
             f'=SUMPRODUCT(({CAT}=$A{rr})*({FFA}<0)*({FFB}<0))', f'=COUNTIFS({CAT},$A{rr},{CHEAP},"Food First")']
    for ci, formula in enumerate(f, 2):
        cell = ca.cell(rr, ci)
        if formula is None:
            ref = f"{get_column_letter(ci)}{rr}"; src = FFA if ci == 5 else FFB
            ca[ref] = ArrayFormula(ref, f"=MEDIAN(IF({CAT}=$A{rr},{src}))")
        else:
            cell.value = formula
        cell.font = f_bold if is_all else f_base; cell.alignment = right
        cell.number_format = PCT if ci in (4, 5, 6, 7, 8) else "0"
    for ci in range(1, 13):
        ca.cell(rr, ci).border = border
        if is_all: ca.cell(rr, ci).fill = fill_grey
ca.conditional_formatting.add("D2:H5", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_BLUE)))
ca.conditional_formatting.add("D2:H5", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=PALE_RED)))
ca.cell(7, 1, "Items listed = rows in Comparison Data; items matched = rows with a computed % difference (all rows, since excluded products are not in this workbook).").font = f_note
ca.cell(8, 1, "Category medians use an array formula MEDIAN(IF(...)). Negative % = Food First (or first-named store) cheaper.").font = f_note

# chart feed: item level by category
ca.cell(10, 1, "Chart data (feeds the Charts tab)").font = f_bold
for i, h in enumerate(["Category", "Item", f"FF vs {STORE_A}", f"FF vs {STORE_B}"], 1):
    c = ca.cell(11, i, h); c.font = f_hdr; c.fill = fill_hdr; c.alignment = center; c.border = border
blocks = {}; rr = 12
for cat in CATEGORIES:
    start = rr
    for idx, r in enumerate(ROWS):
        if r[0] != cat: continue
        src = FIRST + idx
        ca.cell(rr, 1, f"='Comparison Data'!A{src}"); ca.cell(rr, 2, f"='Comparison Data'!B{src}")
        ca.cell(rr, 3, f"='Comparison Data'!M{src}").number_format = PCT
        ca.cell(rr, 4, f"='Comparison Data'!N{src}").number_format = PCT
        for ci in range(1, 5): ca.cell(rr, ci).border = border; ca.cell(rr, ci).font = f_base
        rr += 1
    blocks[cat] = (start, rr - 1)
ca.column_dimensions["B"].width = 26
ca.page_setup.orientation = "landscape"; ca.page_setup.fitToWidth = 1; ca.page_setup.fitToHeight = 0
ca.sheet_properties.pageSetUpPr.fitToPage = True

# ------------------------------------------------------------------ Charts
ch = wb.create_sheet("Charts")
ch["A1"] = "CHARTS"; ch["A1"].font = f_title
ch["A2"] = "Bars below zero = Food First cheaper; above zero = Food First more expensive. Charts read live from the other tabs."; ch["A2"].font = f_note

def vals_only(fmt):
    dl = DataLabelList(); dl.showVal = True; dl.showSerName = False; dl.showCatName = False
    dl.showLegendKey = False; dl.showPercent = False; dl.showLeaderLines = False; dl.numFmt = fmt
    return dl
def style(s, hexc):
    s.graphicalProperties = GraphicalProperties(solidFill=hexc); s.graphicalProperties.line.solidFill = hexc
def base(title, horizontal=False):
    c = BarChart(); c.type = "bar" if horizontal else "col"; c.grouping = "clustered"; c.title = title
    c.y_axis.title = "% difference"; c.y_axis.number_format = "0%"; c.y_axis.majorGridlines.spPr = GraphicalProperties(ln=None)
    c.legend.position = "b"; c.gapWidth = 60; c.overlap = -10; c.style = 10
    c.x_axis.delete = False; c.y_axis.delete = False
    return c

cats = Reference(ca, min_col=1, min_row=2, max_row=5)
c1 = base("Mean price difference by category (Food First vs each supermarket)")
sA = Series(Reference(ca, min_col=4, min_row=1, max_row=5), title_from_data=True); sB = Series(Reference(ca, min_col=6, min_row=1, max_row=5), title_from_data=True)
c1.series.append(sA); c1.series.append(sB); c1.set_categories(cats); style(sA, ORANGE); style(sB, AQUA)
c1.dataLabels = vals_only("0%"); c1.height = 9; c1.width = 20; ch.add_chart(c1, "A4")
c2 = base("Median price difference by category (Food First vs each supermarket)")
sA2 = Series(Reference(ca, min_col=5, min_row=1, max_row=5), title_from_data=True); sB2 = Series(Reference(ca, min_col=7, min_row=1, max_row=5), title_from_data=True)
c2.series.append(sA2); c2.series.append(sB2); c2.set_categories(cats); style(sA2, ORANGE); style(sB2, AQUA)
c2.dataLabels = vals_only("0%"); c2.height = 9; c2.width = 20; ch.add_chart(c2, "A23")
c3 = BarChart(); c3.type = "col"; c3.title = "Number of items where each store is cheapest"; c3.y_axis.title = "Items"; c3.legend = None; c3.style = 10; c3.gapWidth = 80
c3.y_axis.majorGridlines.spPr = GraphicalProperties(ln=None); c3.x_axis.delete = False; c3.y_axis.delete = False
s3 = Series(Reference(sm, min_col=2, min_row=X0 + 1, max_row=X0 + 3), title="Items"); c3.series.append(s3)
c3.set_categories(Reference(sm, min_col=1, min_row=X0 + 1, max_row=X0 + 3)); style(s3, BLUE); c3.dataLabels = vals_only("0")
c3.height = 8; c3.width = 14; ch.add_chart(c3, "A42")
anchor = 60
for cat in CATEGORIES:
    a, b = blocks[cat]
    c = base(f"{cat}: Food First price relative to each supermarket (item level)", horizontal=True)
    sA_ = Series(Reference(ca, min_col=3, min_row=a, max_row=b), title=f"FF vs {STORE_A}"); sB_ = Series(Reference(ca, min_col=4, min_row=a, max_row=b), title=f"FF vs {STORE_B}")
    c.series.append(sA_); c.series.append(sB_); c.set_categories(Reference(ca, min_col=2, min_row=a, max_row=b)); style(sA_, ORANGE); style(sB_, AQUA)
    c.x_axis.scaling.orientation = "maxMin"; c.x_axis.tickLblPos = "low"; c.y_axis.crosses = "max"
    n = b - a + 1; c.height = max(7, 0.6 * n + 3); c.width = 24
    ch.add_chart(c, f"A{anchor}"); anchor += int(c.height * 2) + 2
ch.page_setup.orientation = "portrait"; ch.page_setup.fitToWidth = 1; ch.page_setup.fitToHeight = 0
ch.sheet_properties.pageSetUpPr.fitToPage = True

for s_, col in ((ws, BLUE), (sm, NAVY), (ca, ORANGE), (ch, AQUA)):
    s_.sheet_properties.tabColor = col
wb.save(OUT); print("saved", OUT, "X0", X0, blocks)
